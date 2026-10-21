"""Offline test: runs the transform on saved NHL responses (3 teams)."""
import csv
import os
import sys
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
os.environ["NHL_FIXTURES"] = str(HERE / "fixtures")
sys.path.insert(0, str(HERE.parent / "scripts"))
import update_nhl as u  # noqa: E402


def test_all():
    now = datetime(2026, 10, 21, 7, 45, tzinfo=u.EASTERN)
    assert u.season_id(now) == "20262027"
    assert u.season_id(datetime(2027, 3, 1)) == "20262027"
    assert u.season_id(datetime(2027, 8, 1)) == "20272028"

    rows = u.build_rows(now)
    assert [r["Abbrev"] for r in rows] == ["CAR", "UTA", "MTL"]  # sorted by points
    car, uta, mtl = rows
    assert list(car.keys()) == u.COLUMNS
    assert car["GP"] == 11 and car["PTS"] == 17 and car["GF/GP"] == round(42 / 11, 4)
    assert car["CF%"] == 0.59375 and car["PP%"] == 0.25641 and car["PDO"] == 0.97766
    assert car["Streak"] == "W1" and car["Home GF"] == 13 and car["L10 W"] == 8
    assert mtl["Team"] == "Montréal Canadiens" and mtl["CF%"] == 0.55454  # accent joins
    assert uta["SF/GP"] == 30.1 and uta["Div"] == "Central"
    assert all(v is not None for r in rows for v in r.values())

    out = HERE / "output"
    u.write_csv(out / "nhl_teams.csv", rows)
    u.write_xlsx(out / "nhl_teams.xlsx", rows)
    with open(out / "nhl_teams.csv", encoding="utf-8") as f:
        got = list(csv.DictReader(f))
    assert len(got) == 3 and got[0]["Team"] == "Carolina Hurricanes"

    from openpyxl import load_workbook
    ws = load_workbook(out / "nhl_teams.xlsx")["NHL Teams"]
    assert "NHL_Teams" in ws.tables and ws.tables["NHL_Teams"].ref == "A1:AV4"
    assert ws["J2"].number_format == "0.0%" and ws["B2"].value == "CAR"
    print("all checks passed")


if __name__ == "__main__":
    test_all()
