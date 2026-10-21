"""Pull current NHL team stats and write them to data/ as CSV and Excel.

Sources (NHL public feeds, no key needed):
  - api-web.nhle.com/v1/standings/now          record, goals, home/road, last 10
  - api.nhle.com/stats/rest/en/team/summary     shots, special teams, faceoffs
  - api.nhle.com/stats/rest/en/team/percentages Corsi, Fenwick, 5v5 Sh/Sv, PDO

Outputs:
  data/nhl_teams.csv           latest snapshot (what Excel reads)
  data/nhl_teams.xlsx          same data as an Excel table named NHL_Teams
  data/history/YYYY-MM-DD.csv  one dated snapshot per day, for backtesting
"""

import csv
import json
import os
import sys
import time
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
EASTERN = ZoneInfo("America/New_York")

# Set NHL_FIXTURES to a folder of saved responses to run offline (used for testing)
FIXTURES = os.environ.get("NHL_FIXTURES")

COLUMNS = [
    "Team", "Abbrev", "Conf", "Div", "GP", "W", "L", "OTL", "PTS", "P%", "RW", "ROW",
    "GF", "GA", "GD", "GF/GP", "GA/GP", "SF/GP", "SA/GP", "PP%", "PK%", "FO%",
    "CF%", "FF%", "CF% Close", "5v5 Sh%", "5v5 Sv%", "PDO", "OZ Start%",
    "Home GP", "Home W", "Home L", "Home OTL", "Home GF", "Home GA",
    "Road GP", "Road W", "Road L", "Road OTL", "Road GF", "Road GA",
    "L10 W", "L10 L", "L10 OTL", "L10 GF", "L10 GA", "Streak", "Updated",
]
PCT_COLS = {"P%", "PP%", "PK%", "FO%", "CF%", "FF%", "CF% Close", "5v5 Sh%", "5v5 Sv%", "OZ Start%"}
TWO_DP = {"GF/GP", "GA/GP", "SF/GP", "SA/GP"}


def season_id(now):
    start = now.year if now.month >= 8 else now.year - 1
    return f"{start}{start + 1}"


def get_json(name, url, params=None):
    if FIXTURES:
        return json.loads(Path(FIXTURES, f"{name}.json").read_text(encoding="utf-8"))
    if params:
        url = url + "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": "nhl-stats-updater (GitHub Actions)"})
    last_err = None
    for attempt in range(4):
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                return json.load(resp)
        except Exception as err:  # network hiccup or 5xx: wait and retry
            last_err = err
            time.sleep(5 * (attempt + 1))
    raise RuntimeError(f"Could not fetch {url}: {last_err}")


def per_game(n, gp):
    return round(n / gp, 4) if gp else None


def build_rows(now):
    sid = season_id(now)
    cayenne = {"cayenneExp": f"seasonId={sid} and gameTypeId=2", "limit": "-1"}

    standings = get_json("standings", "https://api-web.nhle.com/v1/standings/now")["standings"]
    summary = get_json("summary", "https://api.nhle.com/stats/rest/en/team/summary", cayenne)["data"]
    pct = get_json("percentages", "https://api.nhle.com/stats/rest/en/team/percentages", cayenne)["data"]

    summary_by_team = {r["teamFullName"]: r for r in summary}
    pct_by_team = {r["teamFullName"]: r for r in pct}
    stamp = now.strftime("%Y-%m-%d %H:%M")

    rows = []
    for r in standings:
        team = r["teamName"]["default"]
        s = summary_by_team.get(team, {})
        p = pct_by_team.get(team, {})
        gp = r.get("gamesPlayed")
        rows.append({
            "Team": team,
            "Abbrev": r["teamAbbrev"]["default"],
            "Conf": r.get("conferenceName"),
            "Div": r.get("divisionName"),
            "GP": gp,
            "W": r.get("wins"),
            "L": r.get("losses"),
            "OTL": r.get("otLosses"),
            "PTS": r.get("points"),
            "P%": r.get("pointPctg"),
            "RW": r.get("regulationWins"),
            "ROW": r.get("regulationPlusOtWins"),
            "GF": r.get("goalFor"),
            "GA": r.get("goalAgainst"),
            "GD": r.get("goalDifferential"),
            "GF/GP": per_game(r.get("goalFor", 0), gp),
            "GA/GP": per_game(r.get("goalAgainst", 0), gp),
            "SF/GP": s.get("shotsForPerGame"),
            "SA/GP": s.get("shotsAgainstPerGame"),
            "PP%": s.get("powerPlayPct"),
            "PK%": s.get("penaltyKillPct"),
            "FO%": s.get("faceoffWinPct"),
            "CF%": p.get("satPct"),
            "FF%": p.get("usatPct"),
            "CF% Close": p.get("satPctClose"),
            "5v5 Sh%": p.get("shootingPct5v5"),
            "5v5 Sv%": p.get("savePct5v5"),
            "PDO": p.get("shootingPlusSavePct5v5"),
            "OZ Start%": p.get("zoneStartPct5v5"),
            "Home GP": r.get("homeGamesPlayed"),
            "Home W": r.get("homeWins"),
            "Home L": r.get("homeLosses"),
            "Home OTL": r.get("homeOtLosses"),
            "Home GF": r.get("homeGoalsFor"),
            "Home GA": r.get("homeGoalsAgainst"),
            "Road GP": r.get("roadGamesPlayed"),
            "Road W": r.get("roadWins"),
            "Road L": r.get("roadLosses"),
            "Road OTL": r.get("roadOtLosses"),
            "Road GF": r.get("roadGoalsFor"),
            "Road GA": r.get("roadGoalsAgainst"),
            "L10 W": r.get("l10Wins"),
            "L10 L": r.get("l10Losses"),
            "L10 OTL": r.get("l10OtLosses"),
            "L10 GF": r.get("l10GoalsFor"),
            "L10 GA": r.get("l10GoalsAgainst"),
            "Streak": f"{r.get('streakCode') or ''}{r.get('streakCount') or ''}",
            "Updated": stamp,
        })

    missing = [row["Team"] for row in rows if row["CF%"] is None or row["PP%"] is None]
    if missing and len(missing) < len(rows):
        print(f"Warning: no stats-feed match for {', '.join(missing)}", file=sys.stderr)

    rows.sort(key=lambda x: (x["PTS"] or 0, x["P%"] or 0), reverse=True)
    return rows


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, rows):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "NHL Teams"
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row[c] for c in COLUMNS])
    last = get_column_letter(len(COLUMNS))
    table = Table(displayName="NHL_Teams", ref=f"A1:{last}{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
    ws.add_table(table)
    for i, col in enumerate(COLUMNS, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = {"Team": 24, "Div": 13, "Updated": 17}.get(col, max(7, len(col) + 3))
        fmt = "0.0%" if col in PCT_COLS else "0.00" if col in TWO_DP else "0.000" if col == "PDO" else None
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "B2"
    wb.save(path)


def main():
    now = datetime.now(EASTERN)
    rows = build_rows(now)
    if len(rows) < 30:
        raise SystemExit(f"Only {len(rows)} teams came back; not overwriting the data files.")

    write_csv(DATA / "nhl_teams.csv", rows)
    write_xlsx(DATA / "nhl_teams.xlsx", rows)
    write_csv(DATA / "history" / f"{now:%Y-%m-%d}.csv", rows)
    print(f"Wrote {len(rows)} teams for season {season_id(now)} at {now:%Y-%m-%d %H:%M} ET")


if __name__ == "__main__":
    main()
